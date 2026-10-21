# scripts/convert_guide.py
from __future__ import annotations

import os
import re
import subprocess
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MD_PATH = Path("documentation/INTERVIEW_GUIDE.md")
HTML_PATH = Path("documentation/INTERVIEW_GUIDE.html")
PDF_PATH = Path("documentation/INTERVIEW_GUIDE.pdf")
DOCX_PATH = Path("documentation/INTERVIEW_GUIDE.docx")


def markdown_to_html(md_text: str) -> str:
    """Convert basic markdown to styled HTML suitable for printing to PDF."""
    html_lines = []
    in_code_block = False
    in_table = False
    table_rows = []

    lines = md_text.splitlines()
    for line in lines:
        # Code block toggle
        if line.startswith("```"):
            if in_code_block:
                html_lines.append("</pre>")
                in_code_block = False
            else:
                lang = line.strip("`").strip()
                html_lines.append(f'<pre class="code-block {lang}">')
                in_code_block = True
            continue

        if in_code_block:
            escaped = (
                line.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            html_lines.append(escaped)
            continue

        # Table rows
        if line.startswith("|") and line.endswith("|"):
            parts = [p.strip() for p in line.strip("|").split("|")]
            if all(re.match(r"^:?-+:?$", p) for p in parts if p):
                continue  # Header separator line
            table_rows.append(parts)
            in_table = True
            continue
        elif in_table:
            # End of table
            html_lines.append('<table class="styled-table">')
            for i, r in enumerate(table_rows):
                tag = "th" if i == 0 else "td"
                html_lines.append("<tr>" + "".join(f"<{tag}>{p}</{tag}>" for p in r) + "</tr>")
            html_lines.append("</table>")
            in_table = False
            table_rows = []

        # Horizontal rule
        if line.strip() == "---":
            html_lines.append('<hr class="divider"/>')
            continue

        # Headings
        if line.startswith("# "):
            html_lines.append(f'<h1 class="doc-title">{line[2:].strip()}</h1>')
            continue
        if line.startswith("## "):
            html_lines.append(f'<h2 class="section-title">{line[3:].strip()}</h2>')
            continue
        if line.startswith("### "):
            html_lines.append(f'<h3 class="sub-title">{line[4:].strip()}</h3>')
            continue
        if line.startswith("#### "):
            html_lines.append(f'<h4 class="sub4-title">{line[5:].strip()}</h4>')
            continue

        # Blockquote
        if line.startswith("> "):
            quote_text = line[2:].strip()
            # formatting
            quote_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", quote_text)
            quote_text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", quote_text)
            html_lines.append(f'<div class="callout">{quote_text}</div>')
            continue

        # Bullet list
        if line.strip().startswith("- ") or line.strip().startswith("• "):
            bullet_text = line.strip()[2:].strip()
            bullet_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", bullet_text)
            bullet_text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", bullet_text)
            bullet_text = re.sub(r"`(.*?)`", r"<code>\1</code>", bullet_text)
            html_lines.append(f'<li class="bullet-item">{bullet_text}</li>')
            continue

        # Numbered list
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if num_match:
            item_text = num_match.group(2)
            item_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", item_text)
            item_text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", item_text)
            item_text = re.sub(r"`(.*?)`", r"<code>\1</code>", item_text)
            html_lines.append(f'<li class="num-item"><strong>{num_match.group(1)}.</strong> {item_text}</li>')
            continue

        # Empty line
        if not line.strip():
            html_lines.append('<div class="spacer"></div>')
            continue

        # Regular paragraph
        p_text = line.strip()
        p_text = re.sub(r"\*\*(.*?)\*\*", r"<strong>\1</strong>", p_text)
        p_text = re.sub(r"\*(.*?)\*", r"<em>\1</em>", p_text)
        p_text = re.sub(r"`(.*?)`", r"<code>\1</code>", p_text)
        html_lines.append(f"<p>{p_text}</p>")

    if in_table and table_rows:
        html_lines.append('<table class="styled-table">')
        for i, r in enumerate(table_rows):
            tag = "th" if i == 0 else "td"
            html_lines.append("<tr>" + "".join(f"<{tag}>{p}</{tag}>" for p in r) + "</tr>")
        html_lines.append("</table>")

    body_content = "\n".join(html_lines)

    html_template = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Aanya — Technical Interview & Architecture Guide</title>
<style>
  @page {{
    size: A4;
    margin: 18mm 16mm 18mm 16mm;
    @bottom-center {{
      content: counter(page);
    }}
  }}
  body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
    color: #1a202c;
    line-height: 1.55;
    font-size: 10.5pt;
    margin: 0;
    padding: 0;
  }}
  h1.doc-title {{
    color: #1e3a8a;
    font-size: 20pt;
    border-bottom: 3px solid #3b82f6;
    padding-bottom: 8px;
    margin-top: 0;
    margin-bottom: 16px;
  }}
  h2.section-title {{
    color: #1e40af;
    font-size: 14pt;
    border-bottom: 1.5px solid #93c5fd;
    padding-bottom: 4px;
    margin-top: 22px;
    margin-bottom: 10px;
    page-break-after: avoid;
  }}
  h3.sub-title {{
    color: #1e293b;
    font-size: 11.5pt;
    margin-top: 14px;
    margin-bottom: 6px;
    page-break-after: avoid;
  }}
  h4.sub4-title {{
    color: #334155;
    font-size: 10.5pt;
    margin-top: 10px;
    margin-bottom: 4px;
    page-break-after: avoid;
  }}
  p {{
    margin: 4px 0 8px 0;
  }}
  .callout {{
    background: #f0fdf4;
    border-left: 4px solid #16a34a;
    padding: 10px 14px;
    margin: 10px 0;
    border-radius: 0 4px 4px 0;
    font-size: 10pt;
    color: #14532d;
  }}
  pre.code-block {{
    background: #0f172a;
    color: #f8fafc;
    padding: 12px;
    border-radius: 6px;
    font-family: Consolas, "Courier New", monospace;
    font-size: 8.5pt;
    line-height: 1.35;
    overflow-x: auto;
    margin: 10px 0;
    page-break-inside: avoid;
    white-space: pre-wrap;
  }}
  code {{
    background: #f1f5f9;
    color: #be185d;
    padding: 1px 4px;
    border-radius: 3px;
    font-family: Consolas, "Courier New", monospace;
    font-size: 9pt;
  }}
  pre.code-block code {{
    background: transparent;
    color: inherit;
    padding: 0;
  }}
  table.styled-table {{
    width: 100%;
    border-collapse: collapse;
    margin: 12px 0;
    font-size: 9.5pt;
    page-break-inside: avoid;
  }}
  table.styled-table th {{
    background: #1e3a8a;
    color: #ffffff;
    font-weight: 600;
    text-align: left;
    padding: 7px 10px;
    border: 1px solid #1e3a8a;
  }}
  table.styled-table td {{
    padding: 6px 10px;
    border: 1px solid #cbd5e1;
  }}
  table.styled-table tr:nth-child(even) {{
    background: #f8fafc;
  }}
  li.bullet-item, li.num-item {{
    margin-bottom: 4px;
    margin-left: 18px;
  }}
  hr.divider {{
    border: 0;
    height: 1px;
    background: #cbd5e1;
    margin: 16px 0;
  }}
  .spacer {{
    height: 6px;
  }}
</style>
</head>
<body>
{body_content}
</body>
</html>
"""
    return html_template


def generate_docx(md_text: str, output_path: Path) -> None:
    """Generate a clean, styled Microsoft Word (.docx) document."""
    doc = Document()

    # Page margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    lines = md_text.splitlines()
    in_code_block = False
    code_lines = []

    for line in lines:
        if line.startswith("```"):
            if in_code_block:
                # Add full code block paragraph
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        # Skip horizontal dividers or markdown tables in simple docx renderer
        if line.strip() == "---":
            continue

        # Headings
        if line.startswith("# "):
            h = doc.add_heading(line[2:].strip(), level=1)
            h.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            continue
        if line.startswith("## "):
            h = doc.add_heading(line[3:].strip(), level=2)
            h.style.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
            continue
        if line.startswith("### "):
            h = doc.add_heading(line[4:].strip(), level=3)
            h.style.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
            continue
        if line.startswith("#### "):
            h = doc.add_heading(line[5:].strip(), level=4)
            h.style.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            continue

        # Callout / Quote
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:].strip().replace("**", "").replace("*", ""))
            run.font.italic = True
            run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
            continue

        # Bullet items
        if line.strip().startswith("- ") or line.strip().startswith("• "):
            text = line.strip()[2:].strip()
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_runs(p, text)
            continue

        # Numbered items
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if num_match:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            _add_formatted_runs(p, num_match.group(2))
            continue

        # Regular paragraph
        if line.strip() and not line.startswith("|"):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            _add_formatted_runs(p, line.strip())

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


def _add_formatted_runs(paragraph, text: str):
    """Parse inline **bold** and `code` markers into docx runs."""
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.font.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xBE, 0x18, 0x5D)
        else:
            paragraph.add_run(tok)


def main():
    if not MD_PATH.exists():
        print(f"File not found: {MD_PATH}")
        return

    md_content = MD_PATH.read_text(encoding="utf-8")

    # 1. Generate HTML
    html_content = markdown_to_html(md_content)
    HTML_PATH.write_text(html_content, encoding="utf-8")
    print(f"Generated HTML: {HTML_PATH}")

    # 2. Generate PDF via Chrome Headless
    chrome_exe = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
    if os.path.exists(chrome_exe):
        html_abs = HTML_PATH.resolve()
        pdf_abs = PDF_PATH.resolve()
        cmd = [
            chrome_exe,
            "--headless=new",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_abs}",
            str(html_abs),
        ]
        result = subprocess.run(cmd, capture_output=True, text=True)
        if result.returncode == 0 and pdf_abs.exists():
            print(f"Generated PDF: {pdf_abs} ({pdf_abs.stat().st_size} bytes)")
        else:
            print("Chrome PDF generation error:", result.stderr)
    else:
        print("Chrome not found for PDF export.")

    # 3. Generate DOCX
    generate_docx(md_content, DOCX_PATH)


if __name__ == "__main__":
    main()
